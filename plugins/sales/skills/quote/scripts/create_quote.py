#!/usr/bin/env python3
"""見積書（.xlsx）を作る。税の計算・採番・ファイル名の固定までを行う。

依存: openpyxl（pip install openpyxl）。PDF への変換は LibreOffice（soffice）があれば --pdf で行う。

入力:
  --company  発行元の情報の JSON（references/company-info.example.json を写して埋めたもの）
  --data     見積の内容の JSON（下のスキーマ）
  --out-dir  保存先のフォルダ（クライアント別のフォルダを指定する）

見積の内容（--data）:
{
  "client_company": "株式会社サンプル",        必須。正式名称（「御中」は付けない）
  "client_name": "営業部 山田 太郎 様",        任意
  "subject": "Web サイト制作 御見積",          必須
  "issue_date": "2026-10-05",                 任意。無ければ今日
  "valid_until": "2026-11-04",                任意。無ければ発行日から 30 日
  "quote_number": "Q-20261005-001",           任意。無ければ日付＋連番で採番
  "remarks": "納期: ご発注から 4 週間",          任意。無ければ発行元の既定の支払条件
  "items": [
    {"name": "初期設定", "qty": 1, "unit": "式", "unit_price": 100000},
    {"name": "交通費（実費）", "qty": 1, "unit": "式", "unit_price": 11000, "tax_included": true}
  ]
}
unit_price は税抜で渡す。税込で受け取った行にだけ "tax_included": true を立てる（税の加算を飛ばす）。

採番: 保存先の親フォルダ（--ledger で変更可）の quote-ledger.csv を読み、同じ日の連番を進めて追記する。
同名のファイルがあっても上書きしない（_2, _3 を付けて新しく保存し、警告を出す）。

使い方:
  python3 create_quote.py --company company-info.json --data quote.json --out-dir ./quotes/株式会社サンプル --pdf
  python3 create_quote.py --company company-info.json --data quote.json --dry-run     # 金額の計算だけ確かめる
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import subprocess
import sys
from datetime import date, datetime, timedelta
from decimal import ROUND_CEILING, ROUND_FLOOR, ROUND_HALF_UP, Decimal
from pathlib import Path

ROUNDING = {"floor": ROUND_FLOOR, "round": ROUND_HALF_UP, "ceil": ROUND_CEILING}
PLACEHOLDER = "{{未設定}}"


def yen(n: int) -> str:
    return f"¥{n:,}"


def parse_iso(value: str | None, field: str) -> date | None:
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        sys.exit(f"エラー: {field} は YYYY-MM-DD で指定してください（受け取った値: {value}）")


def jp_date(d: date) -> str:
    return f"{d.year}年{d.month}月{d.day}日"


def load_company(path: Path) -> dict:
    company = json.loads(path.read_text(encoding="utf-8"))
    if not company.get("company_name") or company["company_name"] == PLACEHOLDER:
        sys.exit("エラー: 発行元の情報に company_name がありません。利用者に確認して設定ファイルに書き戻してください")
    unset = [k for k, v in company.items() if v == PLACEHOLDER]
    if unset:
        print(f"注意: 発行元の情報に未設定の項目があります（見積書には載せません）: {', '.join(unset)}", file=sys.stderr)
    return company


def totals(items: list[dict], tax_rate: Decimal, rounding: str) -> dict:
    """税抜の行は小計に税を 1 回だけ掛ける。税込の行は加算を飛ばす。"""
    ex_subtotal = 0
    in_subtotal = 0
    for i, it in enumerate(items, 1):
        for key in ("name", "qty", "unit_price"):
            if key not in it or it[key] in ("", None):
                sys.exit(f"エラー: 明細 {i} 行目に {key} がありません（単価や数量を推測で埋めない）")
        amount = Decimal(str(it["qty"])) * Decimal(str(it["unit_price"]))
        if amount != amount.to_integral_value():
            sys.exit(f"エラー: 明細 {i} 行目の金額が整数になりません（数量 {it['qty']} × 単価 {it['unit_price']}）")
        it["_amount"] = int(amount)
        if it.get("tax_included"):
            in_subtotal += int(amount)
        else:
            ex_subtotal += int(amount)
    tax = int((Decimal(ex_subtotal) * tax_rate).quantize(Decimal("1"), rounding=ROUNDING[rounding]))
    return {"ex_subtotal": ex_subtotal, "in_subtotal": in_subtotal, "tax": tax, "total": ex_subtotal + tax + in_subtotal}


def next_number(ledger: Path, issue: date, prefix: str) -> str:
    stem = f"{prefix}{issue.strftime('%Y%m%d')}-"
    seq = 0
    if ledger.exists():
        with ledger.open(encoding="utf-8-sig", newline="") as f:
            for row in csv.DictReader(f):
                num = row.get("quote_number", "")
                if num.startswith(stem):
                    m = re.search(r"-(\d+)$", num)
                    if m:
                        seq = max(seq, int(m.group(1)))
    return f"{stem}{seq + 1:03d}"


def append_ledger(ledger: Path, row: dict) -> None:
    new = not ledger.exists()
    ledger.parent.mkdir(parents=True, exist_ok=True)
    with ledger.open("a", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["quote_number", "issue_date", "client_company", "subject", "total", "valid_until", "file"])
        if new:
            w.writeheader()
        w.writerow(row)


def safe_name(text: str) -> str:
    return re.sub(r'[\\/:*?"<>|\n\r\t]', "-", text).strip()


def output_path(out_dir: Path, subject: str, client: str, issue: date) -> tuple[Path, bool]:
    base = f"御見積書_{safe_name(subject)}_{safe_name(client)}御中_{issue.strftime('%Y%m%d')}"
    path = out_dir / f"{base}.xlsx"
    clash = path.exists()
    n = 2
    while path.exists():
        path = out_dir / f"{base}_{n}.xlsx"
        n += 1
    return path, clash


def build_workbook(quote: dict, company: dict, t: dict, tax_label: str):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    font_name = company.get("font", "游ゴシック")
    wb = Workbook()
    ws = wb.active
    ws.title = "御見積書"
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 6, "B": 28, "C": 8, "D": 6, "E": 14, "F": 16, "G": 4}.items():
        ws.column_dimensions[col].width = w

    thin = Side(style="thin", color="333333")
    medium = Side(style="medium", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    box_med = Border(left=medium, right=medium, top=medium, bottom=medium)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")
    head_fill = PatternFill("solid", fgColor="2F5496")
    total_fill = PatternFill("solid", fgColor="DCE6F1")

    ws.merge_cells("A1:G2")
    ws["A1"] = "御 見 積 書"
    ws["A1"].font = Font(name=font_name, size=24, bold=True)
    ws["A1"].alignment = center
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 26

    ws["A4"] = f"{quote['client_company']}　御中"
    ws["A4"].font = Font(name=font_name, size=14, bold=True)
    ws.merge_cells("A4:D4")
    if quote.get("client_name"):
        ws["A5"] = quote["client_name"]
        ws["A5"].font = Font(name=font_name, size=12)
        ws.merge_cells("A5:D5")

    meta = [("見積番号", quote["quote_number"]), ("発行日", quote["_issue_label"]), ("有効期限", quote["_valid_label"])]
    for i, (label, value) in enumerate(meta, start=4):
        ws.cell(row=i, column=5, value=label).alignment = right
        ws.cell(row=i, column=5).font = Font(name=font_name, size=10)
        ws.cell(row=i, column=6, value=value).alignment = right
        ws.cell(row=i, column=6).font = Font(name=font_name, size=10, bold=True)

    ws["A7"] = f"件名: {quote['subject']}"
    ws["A7"].font = Font(name=font_name, size=12, bold=True)
    ws.merge_cells("A7:D7")
    ws["A8"] = "下記の通りお見積もり申し上げます。"
    ws["A8"].font = Font(name=font_name, size=10)
    ws.merge_cells("A8:D8")

    ws.merge_cells("A10:B10")
    ws["A10"] = "合計金額（税込）"
    ws["A10"].alignment = center
    ws["A10"].font = Font(name=font_name, size=11, bold=True, color="FFFFFF")
    ws["A10"].fill = head_fill
    ws.merge_cells("A11:B11")
    ws["A11"] = yen(t["total"])
    ws["A11"].alignment = center
    ws["A11"].font = Font(name=font_name, size=20, bold=True)
    ws["A11"].border = box_med
    ws.row_dimensions[11].height = 36

    ws.merge_cells("E10:F10")
    ws["E10"] = company["company_name"]
    ws["E10"].font = Font(name=font_name, size=12, bold=True)
    ws["E10"].alignment = left
    issuer = []
    for key, fmt in (("postal_code", "〒{}"), ("address", "{}"), ("tel", "TEL: {}"), ("email", "Email: {}"), ("invoice_number", "登録番号: {}")):
        v = company.get(key)
        if v and v != PLACEHOLDER:
            issuer.append(fmt.format(v))
    for offset, value in enumerate(issuer):
        row = 11 + offset
        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=6)
        c = ws.cell(row=row, column=5, value=value)
        c.font = Font(name=font_name, size=9)
        c.alignment = left

    seal = company.get("seal_image_path") or ""
    if seal and seal != PLACEHOLDER and Path(seal).expanduser().exists():
        try:
            img = Image(str(Path(seal).expanduser()))
            img.width = img.height = 60
            ws.add_image(img, "G10")
        except Exception as exc:  # 印影が読めなくても見積書は出す
            print(f"注意: 印影の画像を読み込めませんでした（{exc}）", file=sys.stderr)

    header_row = 17
    for i, h in enumerate(["No", "品目", "数量", "単位", "単価", "金額"], start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        c.font = Font(name=font_name, size=10, bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = center
        c.border = box

    row = header_row + 1
    for idx, item in enumerate(quote["items"], start=1):
        name = f"{item['name']}（税込）" if item.get("tax_included") else item["name"]
        values = [idx, name, item["qty"], item.get("unit", "式"), yen(int(Decimal(str(item["unit_price"])))), yen(item["_amount"])]
        for i, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=i, value=v)
            c.font = Font(name=font_name, size=10)
            c.border = box
            c.alignment = left if i == 2 else right if i in (5, 6) else center
        row += 1
    while row < header_row + 1 + max(len(quote["items"]), 5):
        for i in range(1, 7):
            ws.cell(row=row, column=i, value=None).border = box
        row += 1

    summary: list[tuple[str, int]] = []
    if t["ex_subtotal"] > 0 and t["in_subtotal"] > 0:
        summary += [("小計（税抜対象）", t["ex_subtotal"]), (tax_label, t["tax"]), ("税込項目 計", t["in_subtotal"])]
    elif t["in_subtotal"] > 0:
        summary.append(("小計（税込）", t["in_subtotal"]))
    else:
        summary += [("小計", t["ex_subtotal"]), (tax_label, t["tax"])]
    summary.append(("合計", t["total"]))
    for label, value in summary:
        a = ws.cell(row=row, column=5, value=label)
        a.alignment = right
        a.font = Font(name=font_name, size=10, bold=True)
        a.border = box
        c = ws.cell(row=row, column=6, value=yen(value))
        c.alignment = right
        c.font = Font(name=font_name, size=10, bold=(label == "合計"))
        c.border = box
        if label == "合計":
            a.fill = c.fill = total_fill
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="【備考】").font = Font(name=font_name, size=11, bold=True)
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row + 2, end_column=6)
    c = ws.cell(row=row, column=1, value=quote["_remarks"])
    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    c.font = Font(name=font_name, size=10)
    c.border = box
    ws.row_dimensions[row].height = 24
    row += 4

    bank = company.get("bank") or {}
    if isinstance(bank, dict) and any(v and v != PLACEHOLDER for v in bank.values()):
        ws.cell(row=row, column=1, value="【お振込先】").font = Font(name=font_name, size=11, bold=True)
        row += 1
        for line in (
            f"{bank.get('name', '')} {bank.get('branch', '')}",
            f"{bank.get('account_type', '')} {bank.get('account_number', '')}",
            f"口座名義: {bank.get('account_holder', '')}",
        ):
            ws.cell(row=row, column=1, value=line).font = Font(name=font_name, size=10)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            row += 1

    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    for side in ("left", "right", "top", "bottom"):
        setattr(ws.page_margins, side, 0.5)
    return wb


def convert_pdf(xlsx: Path) -> Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        for candidate in ("/opt/homebrew/bin/soffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
            if Path(candidate).exists():
                soffice = candidate
                break
    if not soffice:
        print("注意: LibreOffice（soffice）が見つからないので PDF は作っていません。表計算ソフトで開いて PDF に書き出すか、LibreOffice を入れてください", file=sys.stderr)
        return None
    proc = subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(xlsx.parent), str(xlsx)], capture_output=True, text=True)
    pdf = xlsx.with_suffix(".pdf")
    if proc.returncode != 0 or not pdf.exists():
        print(f"注意: PDF への変換に失敗しました: {proc.stderr.strip()}", file=sys.stderr)
        return None
    return pdf


def main() -> None:
    p = argparse.ArgumentParser(description="見積書（.xlsx）を作る。税の計算・採番・ファイル名の固定まで行う")
    p.add_argument("--company", required=True, help="発行元の情報の JSON")
    p.add_argument("--data", required=True, help="見積の内容の JSON")
    p.add_argument("--out-dir", default="quotes", help="保存先のフォルダ（クライアント別のフォルダ）")
    p.add_argument("--ledger", help="採番の台帳 CSV（既定: 保存先の親フォルダの quote-ledger.csv）")
    p.add_argument("--pdf", action="store_true", help="LibreOffice があれば PDF にも変換する")
    p.add_argument("--dry-run", action="store_true", help="ファイルを作らず、計算結果と予定のファイル名だけを表示する")
    args = p.parse_args()

    company = load_company(Path(args.company).expanduser())
    quote = json.loads(Path(args.data).expanduser().read_text(encoding="utf-8"))
    for key in ("client_company", "subject", "items"):
        if not quote.get(key):
            sys.exit(f"エラー: 見積の内容に {key} がありません")

    issue = parse_iso(quote.get("issue_date"), "issue_date") or date.today()
    valid = parse_iso(quote.get("valid_until"), "valid_until") or issue + timedelta(days=int(company.get("valid_days", 30)))
    tax_rate = Decimal(str(company.get("tax_rate", "0.10")))
    rounding = company.get("tax_rounding", "round")
    if rounding not in ROUNDING:
        sys.exit("エラー: tax_rounding は floor / round / ceil のどれかにしてください")
    t = totals(quote["items"], tax_rate, rounding)
    tax_label = f"消費税({(tax_rate * 100).normalize():f}%)"

    out_dir = Path(args.out_dir).expanduser()
    ledger = Path(args.ledger).expanduser() if args.ledger else out_dir.parent / "quote-ledger.csv"
    number = quote.get("quote_number") or next_number(ledger, issue, company.get("number_prefix", "Q-"))
    quote["quote_number"] = number
    quote["_issue_label"] = jp_date(issue)
    quote["_valid_label"] = jp_date(valid)
    default_remarks = company.get("default_remarks")
    quote["_remarks"] = quote.get("remarks") or (default_remarks if default_remarks and default_remarks != PLACEHOLDER else "")

    path, clash = output_path(out_dir, quote["subject"], quote["client_company"], issue)
    report = {
        "quote_number": number,
        "issue_date": issue.isoformat(),
        "valid_until": valid.isoformat(),
        "subtotal_ex_tax": t["ex_subtotal"],
        "tax": t["tax"],
        "tax_included_items": t["in_subtotal"],
        "total": t["total"],
        "xlsx": str(path),
        "pdf": None,
        "same_name_existed": clash,
    }
    if args.dry_run:
        report["dry_run"] = True
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return

    try:
        wb = build_workbook(quote, company, t, tax_label)
    except ImportError:
        sys.exit("エラー: openpyxl がありません。pip install openpyxl を実行してください")
    out_dir.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    if clash:
        print("警告: 同じ名前のファイルがすでにあったので、上書きせず別名で保存しました。古いほうの扱いを人に確認してください", file=sys.stderr)
    if args.pdf:
        pdf = convert_pdf(path)
        report["pdf"] = str(pdf) if pdf else None
    append_ledger(
        ledger,
        {"quote_number": number, "issue_date": issue.isoformat(), "client_company": quote["client_company"], "subject": quote["subject"], "total": t["total"], "valid_until": valid.isoformat(), "file": str(path)},
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
